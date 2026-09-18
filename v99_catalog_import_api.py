import csv, io, json, re, base64, zipfile, posixpath, xml.etree.ElementTree as ET
from urllib.parse import parse_qs

def _clean(v):
    return str(v or '').strip()

def _num(v, default=0):
    try:
        s=str(v or '').strip().replace('€','').replace(' ','')
        if s.count(',') and s.count('.'): s=s.replace('.','').replace(',','.')
        else: s=s.replace(',','.')
        return float(s)
    except Exception:
        return default

def _bool(v, default=True):
    s=_clean(v).lower()
    if not s:return default
    return s not in ('0','false','nao','não','no','inativo','inactiva','inactivo')

def _norm(s):
    s=_clean(s).lower()
    trans=str.maketrans({'á':'a','à':'a','ã':'a','â':'a','ä':'a','é':'e','ê':'e','è':'e','ë':'e','í':'i','ì':'i','ï':'i','ó':'o','ô':'o','õ':'o','ò':'o','ö':'o','ú':'u','ù':'u','ü':'u','ç':'c','ñ':'n'})
    return re.sub(r'[^a-z0-9]','',s.translate(trans))

ALIASES={
    'sku':{'sku','codigo','codigoartigo','ref','referencia','referenciaartigo','skureferencia','codigoreferencia','referencia sku'},
    'barcode':{'ean','ean13','ean8','barcode','codigodebarras','gtin'},
    'name':{'nome','name','produto','designacao','designacaoproduto','descricao curta','nomeproduto','designacao curta'},
    'brand':{'marca','brand'},
    'category':{'categoria','category'},
    'subcategory':{'subcategoria','subcategory','subfamilia','subfamilias'},
    'family':{'familia','family'},
    'commercialCategory':{'categoriacomercial','categoria comercial','commercialcategory'},
    'commercialSubcategory':{'subcategoriacomercial','subcategoria comercial','commercialsubcategory'},
    'commercialFamily':{'familiacomercial','familia comercial','commercialfamily'},
    'ridaCategory':{'categoriarida','categoria rida','ridacategory'},
    'ridaSubcategory':{'subcategoriarida','subcategoria rida','ridasubcategory'},
    'ridaFamily':{'familiarida','familia rida','ridafamily'},
    'type':{'tipo','type'},
    'price':{'preco','price','pvp','pvpciva','precovenda','precodevenda'},
    'oldPrice':{'precoantigo','oldprice','oldpvp'},
    'description':{'descricao','description','descricaolonga','descricaodetalhada'},
    'image':{'imagem','image','urlimagem','imagemurl','fotografia','foto'},
    'badge':{'badge','etiqueta'},
    'cost':{'custo','cost','precocusto'},
    'vatRate':{'iva','ivarate','taxaiva'},
    'active':{'ativo','active','estado'},
    'stockMin':{'stockmin','stockminimo','minstock','stockminimoartigo'},
    'stockInitial':{'stock','stockinicial','stockinicialfisico','quantidade','quantidadeinicial'}
}
_ALIAS_NORM={field:{_norm(x) for x in keys} for field,keys in ALIASES.items()}

def _match_field(header):
    n=_norm(header)
    if not n:return None
    if n in _ALIAS_NORM['commercialCategory']:return 'commercialCategory'
    if n in _ALIAS_NORM['commercialSubcategory']:return 'commercialSubcategory'
    if n in _ALIAS_NORM['commercialFamily']:return 'commercialFamily'
    if n in _ALIAS_NORM['ridaCategory']:return 'ridaCategory'
    if n in _ALIAS_NORM['ridaSubcategory']:return 'ridaSubcategory'
    if n in _ALIAS_NORM['ridaFamily']:return 'ridaFamily'
    for field,keys in _ALIAS_NORM.items():
        if n in keys:return field
    if n.startswith('sku') or n.endswith('sku'):return 'sku'
    if 'ean' in n or 'barcode' in n or 'codigodebarras' in n:return 'barcode'
    if 'descricao' in n and ('curta' in n or 'produto' in n):return 'name'
    if 'descricao' in n:return 'description'
    if 'familia' in n:return 'family'
    if 'subfamilia' in n or 'subcategoria' in n:return 'subcategory'
    if 'categoria' in n:return 'category'
    if 'marca' in n:return 'brand'
    if 'pvp' in n or 'preco' in n:return 'price'
    if 'custo' in n:return 'cost'
    if 'iva' in n:return 'vatRate'
    if 'stockmin' in n:return 'stockMin'
    if 'stockinicial' in n:return 'stockInitial'
    if n=='stock' or n.startswith('stock') and 'min' not in n:return 'stockInitial'
    return None

def _image_cell_value(cell):
    try:
        v=cell.value
        if isinstance(v,str) and v.strip():return v.strip()
        link=getattr(cell,'hyperlink',None)
        target=getattr(link,'target',None) if link else None
        if target:return str(target).strip()
    except Exception:pass
    return ''

def _row_payload(raw):
    normalized={}
    for k,v in raw.items():
        field=_match_field(k)
        if field and field not in normalized:normalized[field]=v
    out=dict(normalized)
    if 'name' not in out and 'description' in out:out['name']=out['description']
    for k in ('sku','name','brand','category','subcategory','family','type','barcode','description','image','badge','commercialCategory','commercialSubcategory','commercialFamily','ridaCategory','ridaSubcategory','ridaFamily'):
        out[k]=_clean(out.get(k))
    if out.get('commercialCategory'): out['category']=out['commercialCategory']
    if out.get('commercialSubcategory'): out['subcategory']=out['commercialSubcategory']
    if out.get('commercialFamily'): out['family']=out['commercialFamily']
    out['price']=_num(out.get('price'));out['oldPrice']=_num(out.get('oldPrice'));out['cost']=_num(out.get('cost'));out['vatRate']=_num(out.get('vatRate'),23);out['stockMin']=max(0,int(_num(out.get('stockMin'),0)));out['active']=_bool(out.get('active'),True)
    if 'stockInitial' in normalized:
        out['stockInitial']=max(0,int(_num(normalized.get('stockInitial'),0)));out['stockProvided']=True
    else:
        out['stockInitial']=0;out['stockProvided']=False
    return out

def _find_header(ws):
    best=None
    for row_number,row in enumerate(ws.iter_rows(min_row=1,max_row=min(ws.max_row,60),values_only=True),1):
        fields=[_match_field(v) for v in row if v not in (None,'')]
        if 'sku' in fields and ('name' in fields or 'description' in fields):
            score=3+int('brand' in fields)+int('category' in fields)+int('price' in fields)+int('stockInitial' in fields)+int('image' in fields)
            candidate=(score,row_number,row)
            if best is None or candidate[0]>best[0]:best=candidate
    return (best[1],best[2]) if best else None

def _extract_embedded_images(data,ws):
    result={}
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            drawing_target=None
            for rel in getattr(ws,'_rels',[]) or []:
                if str(getattr(rel,'Type','')).endswith('/drawing'):
                    target=str(getattr(rel,'Target','')).replace('\\','/')
                    if target.startswith('/'):
                        drawing_target=target.lstrip('/')
                    elif target.startswith('xl/'):
                        drawing_target=posixpath.normpath(target)
                    else:
                        drawing_target=posixpath.normpath(posixpath.join('xl/worksheets',target))
                    break
            if not drawing_target or drawing_target not in zf.namelist():return result
            drawing_root=ET.fromstring(zf.read(drawing_target))
            rel_path=posixpath.join(posixpath.dirname(drawing_target),'_rels',posixpath.basename(drawing_target)+'.rels')
            if rel_path not in zf.namelist():return result
            rel_root=ET.fromstring(zf.read(rel_path))
            relmap={r.attrib.get('Id'):r.attrib.get('Target','') for r in rel_root}
            ns={'xdr':'http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing','a':'http://schemas.openxmlformats.org/drawingml/2006/main','r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
            for anchor in drawing_root:
                frm=anchor.find('xdr:from',ns)
                if frm is None:continue
                row_el=frm.find('xdr:row',ns)
                pic=anchor.find('.//xdr:pic',ns)
                blip=pic.find('.//a:blip',ns) if pic is not None else None
                if row_el is None or blip is None:continue
                row=int(row_el.text)+1
                rid=blip.attrib.get('{%s}embed'%ns['r']);target=relmap.get(rid)
                if not target:continue
                target=str(target).replace('\\','/')
                if target.startswith('/'):
                    media=posixpath.normpath(target.lstrip('/'))
                elif target.startswith('xl/'):
                    media=posixpath.normpath(target)
                else:
                    media=posixpath.normpath(posixpath.join(posixpath.dirname(drawing_target),target))
                if media not in zf.namelist():continue
                raw=zf.read(media);ext=media.rsplit('.',1)[-1].lower();mime={'jpg':'image/jpeg','jpeg':'image/jpeg','gif':'image/gif','webp':'image/webp','bmp':'image/bmp'}.get(ext,'image/png')
                result[row]=f'data:{mime};base64,'+base64.b64encode(raw).decode('ascii')
    except Exception:
        return result
    return result

def _read_excel(data,filename=''):
    from openpyxl import load_workbook
    wb=load_workbook(io.BytesIO(data),read_only=False,data_only=True)
    candidates=[];preferred=('produtos','catalogo','catálogo','rida','composição skus','composicao skus')
    for ws in wb.worksheets:
        header=_find_header(ws)
        if not header:continue
        header_row,headers=header;images=_extract_embedded_images(data,ws);rows=[]
        image_header_col=next((i+1 for i,h in enumerate(headers) if _match_field(h)=='image'),None)
        for r in range(header_row+1,ws.max_row+1):
            vals=[ws.cell(r,c).value for c in range(1,ws.max_column+1)]
            if not any(v not in (None,'') for v in vals):continue
            raw=dict(zip(headers,vals));item=_row_payload(raw)
            embedded_img=images.get(r,'')
            cell_img=''
            if image_header_col:
                cell_img=_image_cell_value(ws.cell(r,image_header_col))
            if embedded_img:
                item['image']=embedded_img
            elif cell_img:
                item['image']=cell_img
            if not item.get('sku') and not item.get('name'):continue
            rows.append(item)
        if rows:
            title=ws.title.strip().lower();preference=1 if any(p in title for p in preferred) else 0;candidates.append((preference,len(rows),ws.title,rows))
    if not candidates:raise ValueError('Excel: não encontrei uma folha com cabeçalhos de produto contendo SKU e Descrição/Nome.')
    candidates.sort(key=lambda x:(x[0],x[1]),reverse=True);return candidates[0][3]

def _read_file(data,filename):
    if filename.lower().endswith(('.xlsx','.xlsm')):return _read_excel(data,filename)
    text=data.decode('utf-8-sig',errors='replace');sample=text[:4096]
    try:dialect=csv.Sniffer().sniff(sample,delimiters=',;\t')
    except Exception:dialect=csv.excel;dialect.delimiter=';'
    return [_row_payload(r) for r in csv.DictReader(io.StringIO(text),dialect=dialect)]

def _extract_multipart(handler):
    length=int(handler.headers.get('Content-Length','0') or 0);raw=handler.rfile.read(length);ctype=handler.headers.get('Content-Type','');m=re.search(r'boundary=(?:"([^"]+)"|([^;]+))',ctype)
    if not m:raise ValueError('Multipart inválido: boundary em falta')
    boundary=(m.group(1) or m.group(2)).encode()
    for part in raw.split(b'--'+boundary):
        if b'filename=' not in part:continue
        head,sep,body=part.partition(b'\r\n\r\n')
        if not sep:continue
        hm=re.search(br'filename="([^"]*)"',head);filename=(hm.group(1).decode('utf-8','replace') if hm else 'import.csv');return filename,body.rstrip(b'\r\n-')
    raise ValueError('Nenhum ficheiro encontrado')

def _preview(rows):
    def has_image(r):
        v=_clean(r.get('image','')).lower()
        return v.startswith('data:image/') or v.startswith('http://') or v.startswith('https://')
    return {'total':len(rows),'with_images':sum(1 for r in rows if has_image(r)),'with_stock':sum(1 for r in rows if r.get('stockProvided')),'excel_stock_total':sum(int(r.get('stockInitial',0) or 0) for r in rows if r.get('stockProvided')),'without_stock':sum(1 for r in rows if not r.get('stockProvided')),'with_commercial_classification':sum(1 for r in rows if r.get('commercialCategory') or r.get('category')),'with_rida_classification':sum(1 for r in rows if r.get('ridaCategory') or r.get('ridaSubcategory') or r.get('ridaFamily'))}

def _category_slug(name):
    return re.sub(r'-+','-',re.sub(r'[^a-z0-9]+','-',_norm(name))).strip('-')

def _resolve_classification(cur,category_name,subcategory_name,family_name,auto_create=False):
    category_name=_clean(category_name);subcategory_name=_clean(subcategory_name);family_name=_clean(family_name)
    if not category_name and not subcategory_name and not family_name:return (None,None,None)
    if not category_name:raise ValueError('Classificação: Categoria obrigatória quando é indicada uma subcategoria/família.')
    cur.execute("SELECT id FROM catalog_categories WHERE kind='category' AND lower(name)=lower(%s) AND active=true ORDER BY id LIMIT 1",(category_name,))
    rc=cur.fetchone()
    if not rc:
        if not auto_create: raise ValueError('Categoria não encontrada: '+category_name)
        cur.execute("INSERT INTO catalog_categories(name,slug,parent_id,kind,active) VALUES(%s,%s,NULL,'category',true) RETURNING id",(category_name,_category_slug(category_name)))
        cid=cur.fetchone()[0]
    else: cid=rc[0]
    sid=fid=None
    if subcategory_name:
        cur.execute("SELECT id FROM catalog_categories WHERE kind='subcategory' AND lower(name)=lower(%s) AND parent_id=%s AND active=true LIMIT 1",(subcategory_name,cid))
        rs=cur.fetchone()
        if not rs:
            if not auto_create: raise ValueError('Subcategoria não encontrada: '+subcategory_name+' em '+category_name)
            cur.execute("INSERT INTO catalog_categories(name,slug,parent_id,kind,active) VALUES(%s,%s,%s,'subcategory',true) RETURNING id",(subcategory_name,_category_slug(subcategory_name),cid))
            sid=cur.fetchone()[0]
        else: sid=rs[0]
    if family_name:
        if not sid:raise ValueError('Família indicada sem subcategoria: '+family_name)
        cur.execute("SELECT id FROM catalog_categories WHERE kind='family' AND lower(name)=lower(%s) AND parent_id=%s AND active=true LIMIT 1",(family_name,sid))
        rf=cur.fetchone()
        if not rf:
            if not auto_create: raise ValueError('Família não encontrada: '+family_name+' em '+subcategory_name)
            cur.execute("INSERT INTO catalog_categories(name,slug,parent_id,kind,active) VALUES(%s,%s,%s,'family',true) RETURNING id",(family_name,_category_slug(family_name),sid))
            fid=cur.fetchone()[0]
        else: fid=rf[0]
    return cid,sid,fid

def _save_classifications(cur,pid,row):
    from v9_10_13_product_classification_api import save_one
    cc,cs,cf=_resolve_classification(cur,row.get('commercialCategory') or row.get('category'),row.get('commercialSubcategory') or row.get('subcategory'),row.get('commercialFamily') or row.get('family'),auto_create=True)
    rc,rs,rf=_resolve_classification(cur,row.get('ridaCategory'),row.get('ridaSubcategory'),row.get('ridaFamily'))
    if cc or cs or cf:
        save_one(cur,pid,'commercial',{'categoryId':cc,'subcategoryId':cs,'familyId':cf})
    if rc or rs or rf:
        if (row.get('ridaCategory') or '').strip().lower()!='rida':
            raise ValueError('Classificação RIDA: a Categoria RIDA deve ser exatamente "RIDA".')
        if (row.get('ridaSubcategory') or '').strip().lower()!='máquinas a bateria':
            raise ValueError('Classificação RIDA: a Subcategoria RIDA deve ser exatamente "Máquinas a bateria".')
        if (row.get('ridaFamily') or '').strip().lower() not in ('construção','jardim','acessórios'):
            raise ValueError('Classificação RIDA: a Família RIDA deve ser Construção, Jardim ou Acessórios.')
        save_one(cur,pid,'rida',{'categoryId':rc,'subcategoryId':rs,'familyId':rf})

def _import_rows(rows):
    from v98_catalog_products_api import handle_post
    import psycopg,os
    if not os.environ.get('DATABASE_URL'):raise RuntimeError('DATABASE_URL não configurado')
    errors=[];created=updated=images=stock_added=0
    with psycopg.connect(os.environ['DATABASE_URL']) as conn:
        with conn.cursor() as cur:
            for line,row in enumerate(rows,2):
                cur.execute('SELECT stock FROM mm_product_stock WHERE sku=%s',(row['sku'],));stock_row=cur.fetchone();had_stock=stock_row is not None
                cur.execute('SAVEPOINT import_row')
                try:
                    payload={k:v for k,v in row.items() if k not in ('stockInitial','stockProvided','_filename','commercialCategory','commercialSubcategory','commercialFamily','ridaCategory','ridaSubcategory','ridaFamily')}
                    # "MANTER_IMAGEM_ATUAL" (ou vazio) significa preservar exatamente a imagem atual.
                    image_value=_clean(payload.get('image',''))
                    if not image_value.lower().startswith(('data:image/','http://','https://')):
                        payload.pop('image',None)
                    else:
                        images+=1
                    captured=[]
                    handle_post('/api/catalog/products',payload,lambda status,payload:captured.append((status,payload)))
                    status,result=captured[-1] if captured else (500,{'ok':False,'error':'Sem resposta'})
                    if status>=300 or not result.get('ok'):
                        raise RuntimeError(result.get('error','Erro desconhecido'))
                    _save_classifications(cur,int(result.get('id')),row)
                    if had_stock:updated+=1
                    else:created+=1
                except Exception as row_error:
                    cur.execute('ROLLBACK TO SAVEPOINT import_row')
                    errors.append({'linha':line,'sku':row['sku'],'error':str(row_error)})
                    cur.execute('RELEASE SAVEPOINT import_row')
                    continue
                cur.execute('RELEASE SAVEPOINT import_row')
                if row.get('stockProvided'):
                    initial=int(row.get('stockInitial',0) or 0)
                    cur.execute('SELECT stock FROM mm_product_stock WHERE sku=%s FOR UPDATE',(row['sku'],))
                    sr=cur.fetchone()
                    current=int(sr[0]) if sr else 0
                    if sr:
                        cur.execute('UPDATE mm_product_stock SET stock=%s,stock_min=%s,updated_at=NOW() WHERE sku=%s',(initial,int(row.get('stockMin',0) or 0),row['sku']))
                    else:
                        cur.execute('INSERT INTO mm_product_stock(sku,stock,stock_min) VALUES(%s,%s,%s)',(row['sku'],initial,int(row.get('stockMin',0) or 0)))
                    delta=initial-current
                    if delta:
                        movement_type='entrada' if delta>0 else 'ajuste'
                        cur.execute("INSERT INTO mm_stock_movements(sku,movement_type,delta,resulting_stock,reason,notes,created_by) VALUES(%s,%s,%s,%s,%s,%s,%s)",(row['sku'],movement_type,delta,initial,'Importação de catálogo','Stock definido pelo Excel '+str(row.get('_filename','')),'importador-excel'))
                        stock_added+=delta
        conn.commit()
    return {'created':created,'updated':updated,'images':images,'stock_added':stock_added,'errors':errors}

def handle_upload(handler,send_json):
    try:
        filename,data=_extract_multipart(handler)
        if not filename.lower().endswith(('.csv','.xlsx','.xlsm')):raise ValueError('Formato não suportado. Use CSV ou Excel (.xlsx).')
        rows=_read_file(data,filename)
        if not rows:raise ValueError('O ficheiro não contém linhas de dados.')
        if len(rows)>5000:raise ValueError('Limite de 5000 linhas por importação.')
        bad=[i+2 for i,r in enumerate(rows) if not r.get('sku') or not r.get('name')]
        if bad:raise ValueError('SKU e Nome são obrigatórios. Linhas inválidas: '+', '.join(map(str,bad[:20]))+('…' if len(bad)>20 else ''))
        qs=parse_qs(handler.path.split('?',1)[1] if '?' in handler.path else '');mode=(qs.get('mode') or [''])[0]
        if mode=='preview':send_json(200,{'ok':True,'filename':filename,'preview':_preview(rows)});return True
        for r in rows:r['_filename']=filename
        result=_import_rows(rows);send_json(200,{'ok':not result['errors'],'filename':filename,'total':len(rows),'created':result['created'],'updated':result['updated'],'images':result['images'],'stockAdded':result['stock_added'],'failed':len(result['errors']),'errors':result['errors'][:50],'message':f"Importação concluída: {result['created']} novos, {result['updated']} atualizados, {result['images']} imagens e {result['stock_added']} unidades de stock aplicadas do Excel."});return True
    except ValueError as e:send_json(400,{'ok':False,'error':str(e)});return True
    except Exception as e:send_json(503,{'ok':False,'error':f'Importação: {e}'});return True